#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Render the daily-report markdown into a styled .docx + programmatic QC.

Follows the colleague layout spec: A4, Noto Sans CJK SC, centered title,
4-column stock table with bold header repeated across pages, compact body.
QC reopens the file and checks structure (title/sections/table) — the
render-PNG visual pass is an agent-level step documented in SKILL.md.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def _set_font(run, font: str, size: int, bold: bool = False) -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def _repeat_header(row) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    el = trPr.makeelement(qn("w:tblHeader"), {qn("w:val"): "true"})
    trPr.append(el)


def markdown_to_docx(md_text: str, out_path: Path, *, font: str = "Noto Sans CJK SC") -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)  # A4

    lines = md_text.replace("\r\n", "\n").splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1
            continue
        if ln.startswith("# "):  # report title
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_font(p.add_run(ln[2:].strip()), font, 18, bold=True)
            i += 1
            continue
        if ln.startswith("## "):
            p = doc.add_paragraph()
            _set_font(p.add_run(ln[3:].strip()), font, 14, bold=True)
            i += 1
            continue
        if ln.startswith("### "):
            p = doc.add_paragraph()
            _set_font(p.add_run(ln[4:].strip()), font, 12, bold=True)
            i += 1
            continue
        if ln.strip().startswith("|"):
            # collect the table block
            block: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = lines[i].strip().strip("|")
                cells = [c.strip() for c in raw.split("|")]
                if not re.match(r"^[\s:\-]+$", raw.replace("|", "")):
                    block.append(cells)
                i += 1
            if block:
                ncols = max(len(r) for r in block)
                table = doc.add_table(rows=len(block), cols=ncols)
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, row_cells in enumerate(block):
                    for c_idx in range(ncols):
                        cell = table.cell(r_idx, c_idx)
                        cell.text = ""
                        run = cell.paragraphs[0].add_run(
                            row_cells[c_idx] if c_idx < len(row_cells) else ""
                        )
                        _set_font(run, font, 10.5, bold=(r_idx == 0))
                    if r_idx == 0:
                        _repeat_header(table.rows[0])
            continue
        # plain paragraph (bold markers ** ** stripped to plain text runs)
        p = doc.add_paragraph()
        plain = re.sub(r"\*\*(.+?)\*\*", r"\1", ln)
        _set_font(p.add_run(plain), font, 11)
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def qc_docx(path: Path) -> dict:
    """Programmatic QC: reopen and verify structure per the layout spec."""
    errors: list[str] = []
    doc = Document(str(path))
    paras = [p for p in doc.paragraphs if p.text.strip()]
    if not paras:
        errors.append("document has no paragraphs")
    if not any(p.text.startswith("土耳其股市热点日报") for p in paras):
        errors.append("title paragraph missing")
    if not any("一、热点资讯" in p.text for p in paras):
        errors.append("热点资讯 section missing")
    if not any("二、个股分析" in p.text for p in paras):
        errors.append("个股分析 section missing")
    tables = doc.tables
    stock_table = tables[-1] if tables else None
    if stock_table is None:
        errors.append("stock table missing")
    else:
        if len(stock_table.columns) != 4:
            errors.append(f"stock table has {len(stock_table.columns)} columns (expect 4)")
        if not all(c.text.strip() for c in stock_table.rows[0].cells):
            errors.append("stock table header has empty cells")
    title_p = next((p for p in paras if p.text.startswith("土耳其股市热点日报")), None)
    if title_p is not None and title_p.alignment != WD_ALIGN_PARAGRAPH.CENTER:
        errors.append("title not centered")
    return {"ok": len(errors) == 0, "errors": errors,
            "paragraphs": len(paras), "tables": len(tables),
            "stock_rows": (len(stock_table.rows) - 1) if stock_table else 0}
